"""
frontend/export.py

In-memory export builders for SkadVault AI conversation results.

Public API:
    build_pdf(messages, scope, tenant_id)  -> bytes
    build_docx(messages, scope, tenant_id) -> bytes
    build_csv(messages, scope, tenant_id)  -> bytes

scope: "last" | "full"
"""

import csv
import io
import re
from datetime import datetime


# ─────────────────────────────────────────────────────────────────────────────
# Shared helpers
# ─────────────────────────────────────────────────────────────────────────────

def _get_qa_pairs(messages: list, scope: str) -> list[dict]:
    """Return a list of {question, answer, sources} dicts for the chosen scope."""
    if scope == "last":
        for i in range(len(messages) - 1, -1, -1):
            if messages[i]["role"] == "assistant":
                question = ""
                for j in range(i - 1, -1, -1):
                    if messages[j]["role"] == "user":
                        question = messages[j]["content"]
                        break
                return [{
                    "question": question,
                    "answer": messages[i]["content"],
                    "sources": messages[i].get("sources", []),
                }]
        return []

    # scope == "full"
    pairs: list[dict] = []
    i = 0
    while i < len(messages):
        if messages[i]["role"] == "user":
            question = messages[i]["content"]
            if i + 1 < len(messages) and messages[i + 1]["role"] == "assistant":
                pairs.append({
                    "question": question,
                    "answer": messages[i + 1]["content"],
                    "sources": messages[i + 1].get("sources", []),
                })
                i += 2
                continue
        i += 1
    return pairs


def _source_meta(src: dict) -> tuple[str, str]:
    """Returns (type_str, location_str) for a source chunk."""
    meta = src.get("metadata", {})
    stype = meta.get("source_type", "").upper()
    parts: list[str] = []
    if meta.get("page") is not None:
        parts.append(f"Page {meta['page']}")
    if meta.get("table"):
        parts.append(f"Table: {meta['table']}")
    if meta.get("sheet"):
        parts.append(f"Sheet: {meta['sheet']}")
    return stype, "  ·  ".join(parts)


def _strip_md(text: str) -> str:
    """Convert basic Markdown to plain text for CSV."""
    text = re.sub(r'\*\*(.+?)\*\*', r'\1', text)
    text = re.sub(r'\*(.+?)\*', r'\1', text)
    text = re.sub(r'^#{1,6}\s+', '', text, flags=re.MULTILINE)
    return text.strip()


def _safe_latin1(s: str) -> str:
    """Replace characters outside Latin-1 so ReportLab built-in fonts don't error."""
    return s.encode("latin-1", errors="replace").decode("latin-1")


# ─────────────────────────────────────────────────────────────────────────────
# PDF builder (ReportLab)
# ─────────────────────────────────────────────────────────────────────────────

def build_pdf(messages: list, scope: str, tenant_id: str) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.platypus import (
        HRFlowable, Paragraph, SimpleDocTemplate, Spacer,
    )

    pairs = _get_qa_pairs(messages, scope)
    if not pairs:
        raise ValueError("No content to export.")

    ts = datetime.now().strftime("%d %B %Y, %H:%M")
    buf = io.BytesIO()

    doc = SimpleDocTemplate(
        buf,
        pagesize=A4,
        leftMargin=2.5 * cm,
        rightMargin=2.5 * cm,
        topMargin=2.5 * cm,
        bottomMargin=2.5 * cm,
        title="SkadVault AI Report",
        author="SkadVault AI",
    )

    base = getSampleStyleSheet()

    brand = ParagraphStyle(
        "Brand",
        parent=base["Heading1"],
        fontSize=20,
        textColor=colors.HexColor("#1a1a2e"),
        spaceAfter=4,
    )
    meta_style = ParagraphStyle(
        "Meta",
        parent=base["Normal"],
        fontSize=9,
        textColor=colors.grey,
        spaceAfter=14,
    )
    q_style = ParagraphStyle(
        "Q",
        parent=base["Normal"],
        fontSize=11,
        fontName="Helvetica-Bold",
        textColor=colors.HexColor("#1a1a2e"),
        spaceBefore=14,
        spaceAfter=6,
    )
    body = ParagraphStyle(
        "Body",
        parent=base["Normal"],
        fontSize=10,
        leading=15,
        spaceAfter=4,
    )
    bullet = ParagraphStyle(
        "Bullet",
        parent=base["Normal"],
        fontSize=10,
        leading=15,
        leftIndent=14,
        spaceAfter=2,
    )
    h2 = ParagraphStyle(
        "H2",
        parent=base["Heading2"],
        fontSize=12,
        spaceBefore=8,
        spaceAfter=4,
    )
    src_heading = ParagraphStyle(
        "SrcHeading",
        parent=base["Normal"],
        fontSize=9,
        fontName="Helvetica-Bold",
        textColor=colors.grey,
        spaceBefore=8,
        spaceAfter=2,
    )
    src_item = ParagraphStyle(
        "SrcItem",
        parent=base["Normal"],
        fontSize=9,
        leftIndent=12,
        spaceAfter=2,
        textColor=colors.HexColor("#555555"),
    )

    story = []

    # Report header
    story.append(Paragraph(_safe_latin1("SkadVault AI"), brand))
    story.append(Paragraph(
        _safe_latin1(f"Workspace: {tenant_id}  ·  Generated: {ts}"),
        meta_style,
    ))
    story.append(HRFlowable(width="100%", thickness=1, color=colors.HexColor("#dddddd")))
    story.append(Spacer(1, 10))

    for idx, pair in enumerate(pairs, 1):
        # Question
        story.append(Paragraph(
            _safe_latin1(f"Q{idx}. {pair['question']}"),
            q_style,
        ))

        # Answer — parse Markdown line by line
        for line in pair["answer"].split("\n"):
            stripped = line.strip()
            if not stripped:
                story.append(Spacer(1, 3))
                continue
            if stripped.startswith("# "):
                story.append(Paragraph(_safe_latin1(stripped[2:]), h2))
            elif stripped.startswith("## "):
                story.append(Paragraph(_safe_latin1(stripped[3:]), h2))
            elif stripped.startswith("### "):
                story.append(Paragraph(_safe_latin1(stripped[4:]), h2))
            elif re.match(r"^[-*]\s", stripped):
                content = _safe_latin1(stripped[2:])
                content = re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", content)
                story.append(Paragraph(f"• {content}", bullet))
            elif re.match(r"^\d+\.\s", stripped):
                content = re.sub(r"^\d+\.\s", "", stripped)
                content = re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", _safe_latin1(content))
                story.append(Paragraph(content, bullet))
            else:
                html = re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", _safe_latin1(stripped))
                html = re.sub(r"\*(.+?)\*", r"<i>\1</i>", html)
                story.append(Paragraph(html, body))

        # Sources
        sources = pair.get("sources", [])
        if sources:
            story.append(Paragraph("Sources", src_heading))
            for src in sources:
                stype, location = _source_meta(src)
                label = "  ·  ".join(p for p in [stype, location] if p) or "Source"
                story.append(Paragraph(f"<i>{_safe_latin1(label)}</i>", src_item))

        if idx < len(pairs):
            story.append(Spacer(1, 10))
            story.append(HRFlowable(
                width="100%", thickness=0.5, color=colors.HexColor("#eeeeee"),
            ))

    doc.build(story)
    return buf.getvalue()


# ─────────────────────────────────────────────────────────────────────────────
# Word builder (python-docx)
# ─────────────────────────────────────────────────────────────────────────────

def build_docx(messages: list, scope: str, tenant_id: str) -> bytes:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Cm, Pt, RGBColor

    pairs = _get_qa_pairs(messages, scope)
    if not pairs:
        raise ValueError("No content to export.")

    ts = datetime.now().strftime("%d %B %Y, %H:%M")
    doc = Document()

    for section in doc.sections:
        section.top_margin = Cm(2.5)
        section.bottom_margin = Cm(2.5)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)

    # Title
    title = doc.add_heading("SkadVault AI", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for run in title.runs:
        run.font.color.rgb = RGBColor(0x1A, 0x1A, 0x2E)

    meta = doc.add_paragraph()
    meta_run = meta.add_run(f"Workspace: {tenant_id}  ·  Generated: {ts}")
    meta_run.font.size = Pt(9)
    meta_run.font.color.rgb = RGBColor(0x88, 0x88, 0x88)

    doc.add_paragraph()

    for idx, pair in enumerate(pairs, 1):
        # Question heading
        q_head = doc.add_heading(f"Q{idx}. {pair['question']}", level=2)
        for run in q_head.runs:
            run.font.color.rgb = RGBColor(0x1A, 0x1A, 0x2E)

        # Answer — parse Markdown
        _add_md_to_docx(doc, pair["answer"])

        # Sources
        sources = pair.get("sources", [])
        if sources:
            src_label = doc.add_paragraph()
            src_run = src_label.add_run("Sources")
            src_run.bold = True
            src_run.font.size = Pt(9)
            src_run.font.color.rgb = RGBColor(0x88, 0x88, 0x88)

            for src in sources:
                stype, location = _source_meta(src)
                label = "  ·  ".join(p for p in [stype, location] if p) or "Source"
                sp = doc.add_paragraph(label, style="List Bullet")
                for run in sp.runs:
                    run.font.size = Pt(9)
                    run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

        if idx < len(pairs):
            doc.add_paragraph()

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _add_md_to_docx(doc, text: str) -> None:
    """Parse basic Markdown and append paragraphs to a python-docx Document."""
    for line in text.split("\n"):
        stripped = line.strip()
        if not stripped:
            doc.add_paragraph()
            continue
        if stripped.startswith("# "):
            doc.add_heading(stripped[2:], level=1)
        elif stripped.startswith("## "):
            doc.add_heading(stripped[3:], level=2)
        elif stripped.startswith("### "):
            doc.add_heading(stripped[4:], level=3)
        elif re.match(r"^[-*]\s", stripped):
            p = doc.add_paragraph(style="List Bullet")
            _add_md_runs(p, stripped[2:])
        elif re.match(r"^\d+\.\s", stripped):
            content = re.sub(r"^\d+\.\s", "", stripped)
            p = doc.add_paragraph(style="List Number")
            _add_md_runs(p, content)
        else:
            p = doc.add_paragraph()
            _add_md_runs(p, stripped)


def _add_md_runs(para, text: str) -> None:
    """Add runs to a paragraph, honouring **bold** and *italic* markers."""
    # Split on bold spans first
    parts = re.split(r"(\*\*[^*]+\*\*)", text)
    for part in parts:
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            para.add_run(part[2:-2]).bold = True
        else:
            # Handle *italic* within the remaining text
            italic_parts = re.split(r"(\*[^*]+\*)", part)
            for ip in italic_parts:
                if ip.startswith("*") and ip.endswith("*") and len(ip) > 2:
                    para.add_run(ip[1:-1]).italic = True
                elif ip:
                    para.add_run(ip)


# ─────────────────────────────────────────────────────────────────────────────
# CSV builder
# ─────────────────────────────────────────────────────────────────────────────

def build_csv(messages: list, scope: str, tenant_id: str) -> bytes:
    pairs = _get_qa_pairs(messages, scope)
    if not pairs:
        raise ValueError("No content to export.")

    ts = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow([
        "Question", "Answer", "Source Type", "Source Location",
        "Workspace", "Generated At",
    ])

    for pair in pairs:
        plain_answer = _strip_md(pair["answer"])
        sources = pair.get("sources", [])
        if sources:
            for src in sources:
                stype, location = _source_meta(src)
                writer.writerow([
                    pair["question"],
                    plain_answer,
                    stype,
                    location,
                    tenant_id,
                    ts,
                ])
        else:
            writer.writerow([
                pair["question"],
                plain_answer,
                "",
                "",
                tenant_id,
                ts,
            ])

    # utf-8-sig writes the BOM so Excel opens the file correctly
    return buf.getvalue().encode("utf-8-sig")
